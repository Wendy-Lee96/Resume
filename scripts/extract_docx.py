import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)

out = []
out.append("===== PARAGRAPHS =====")
for p in doc.paragraphs:
    style = p.style.name if p.style else ""
    txt = p.text.strip()
    if txt:
        out.append(f"[{style}] {txt}")

out.append("\n===== TABLES =====")
for ti, table in enumerate(doc.tables):
    out.append(f"\n--- TABLE {ti} ---")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        out.append(" | ".join(cells))

text = "\n".join(out)
with open("resume_extract.txt", "w", encoding="utf-8") as f:
    f.write(text)
print(text)
